import os
import shutil
from uuid import uuid4

from docx import Document as DocxDocument
from fastapi import UploadFile
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.document import Document
from app.models.document_chunk import DocumentChunk
from app.models.hospital import Hospital
from app.models.rag_passage import RAGPassage


ALLOWED_EXTENSIONS = [".pdf", ".docx", ".txt"]


def is_allowed_file(filename: str) -> bool:
    _, extension = os.path.splitext(filename.lower())
    return extension in ALLOWED_EXTENSIONS


def extract_text_from_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
        return file.read()


def extract_text_from_pdf(file_path: str) -> str:
    text_parts = []

    reader = PdfReader(file_path)

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text_parts.append(page_text)

    return "\n".join(text_parts)


def extract_text_from_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)

    paragraphs = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text.strip())

    return "\n".join(paragraphs)


def extract_text(file_path: str) -> str:
    _, extension = os.path.splitext(file_path.lower())

    if extension == ".txt":
        return extract_text_from_txt(file_path)

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    raise ValueError("Unsupported file type")


def split_text_into_chunks(
    text: str,
    chunk_size: int = 700,
    overlap: int = 100
):
    cleaned_text = " ".join(text.split())

    if not cleaned_text:
        return []

    chunks = []
    start = 0

    while start < len(cleaned_text):
        end = start + chunk_size
        chunk = cleaned_text[start:end]

        if chunk.strip():
            chunks.append(chunk.strip())

        start = end - overlap

        if start < 0:
            start = 0

        if start >= len(cleaned_text):
            break

    return chunks


def save_document_chunks(
    db: Session,
    hospital_id: int,
    document_id: int,
    chunks: list[str]
):
    saved_chunks = []

    for index, chunk in enumerate(chunks):
        new_chunk = DocumentChunk(
            hospital_id=hospital_id,
            document_id=document_id,
            chunk_index=index,
            content=chunk
        )

        db.add(new_chunk)
        saved_chunks.append(new_chunk)

    db.commit()

    for chunk in saved_chunks:
        db.refresh(chunk)

    return saved_chunks


def upload_document(
    db: Session,
    hospital_id: int,
    uploaded_by_id: int,
    file: UploadFile,
    document_type: str | None = None
):
    hospital = db.query(Hospital).filter(Hospital.id == hospital_id).first()

    if not hospital:
        raise ValueError("Hospital not found")

    if not file.filename:
        raise ValueError("File name is required")

    if not is_allowed_file(file.filename):
        raise ValueError("Only PDF, DOCX, and TXT files are allowed")

    upload_folder = os.path.join(
        "app",
        "uploads",
        f"hospital_{hospital_id}"
    )

    os.makedirs(upload_folder, exist_ok=True)

    _, extension = os.path.splitext(file.filename)
    stored_file_name = f"{uuid4()}{extension.lower()}"

    file_path = os.path.join(upload_folder, stored_file_name)

    max_bytes = int(settings.MAX_UPLOAD_MB) * 1024 * 1024
    written = 0

    with open(file_path, "wb") as buffer:
        while True:
            block = file.file.read(1024 * 1024)
            if not block:
                break
            written += len(block)
            if written > max_bytes:
                buffer.close()
                os.remove(file_path)
                raise ValueError(
                    f"File is too large. The maximum size is {settings.MAX_UPLOAD_MB} MB."
                )
            buffer.write(block)

    extracted_text = extract_text(file_path)

    if not extracted_text.strip():
        if os.path.exists(file_path):
            os.remove(file_path)

        raise ValueError("No readable text found in this document")

    new_document = Document(
        hospital_id=hospital_id,
        uploaded_by_id=uploaded_by_id,
        file_name=file.filename,
        stored_file_name=stored_file_name,
        file_path=file_path,
        document_type=document_type
    )

    db.add(new_document)
    db.commit()
    db.refresh(new_document)

    chunks = split_text_into_chunks(extracted_text)

    save_document_chunks(
        db=db,
        hospital_id=hospital_id,
        document_id=new_document.id,
        chunks=chunks
    )

    return new_document


def get_all_documents(db: Session, hospital_id: int | None = None):
    query = db.query(Document)

    if hospital_id:
        query = query.filter(Document.hospital_id == hospital_id)

    return query.order_by(Document.created_at.desc()).all()


def get_document_by_id(db: Session, document_id: int):
    return db.query(Document).filter(Document.id == document_id).first()


def get_chunks_by_document_id(db: Session, document_id: int):
    return (
        db.query(DocumentChunk)
        .filter(DocumentChunk.document_id == document_id)
        .order_by(DocumentChunk.chunk_index.asc())
        .all()
    )


def delete_document(db: Session, document_id: int):
    document = get_document_by_id(db, document_id)

    if not document:
        return None

    # Remove everything derived from the document, otherwise the chatbot keeps
    # answering from facts in a file the admin has already deleted.
    db.query(RAGPassage).filter(
        RAGPassage.document_id == document_id
    ).delete(synchronize_session=False)

    db.query(DocumentChunk).filter(
        DocumentChunk.document_id == document_id
    ).delete()

    if os.path.exists(document.file_path):
        os.remove(document.file_path)

    db.delete(document)
    db.commit()

    return document